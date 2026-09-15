# -*- coding: utf-8 -*-
"""提取PPTX全部文本，用于了解被测项目背景。"""
import sys
from pptx import Presentation
from pptx.util import Emu

path = sys.argv[1]
prs = Presentation(path)
print(f"幻灯片尺寸: {prs.slide_width} x {prs.slide_height} EMU = {Emu(prs.slide_width).inches:.2f} x {Emu(prs.slide_height).inches:.2f} in")
print(f"页数: {len(prs.slides)}")
for i, slide in enumerate(prs.slides, 1):
    print(f"\n===== Slide {i} =====")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs)
                if text.strip():
                    print(f"  [{shape.shape_type}] {text}")
        if shape.shape_type == 13:  # picture
            print(f"  [图片] {shape.name}")
        if shape.has_table:
            tbl = shape.table
            for row in tbl.rows:
                cells = [c.text.strip() for c in row.cells]
                print("  [表] " + " | ".join(cells))
